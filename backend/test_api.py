"""Security boundaries and complete authored-demo learning flow."""
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient

from backend.app import COOKIE, create_app


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    return create_app(tmp_path / "test.sqlite3")


@pytest.fixture
def client(app):
    with TestClient(app) as browser:
        assert browser.post("/api/session", json={"name": "Ada"}).status_code == 200
        yield browser


def demo(client):
    return client.get("/api/bootstrap").json()["courses"][0]


def new_course(client):
    response = client.post("/api/courses", json={"name": "My course", "code": "CS101", "examDate": "2026-12-01"})
    assert response.status_code == 200
    return response.json()


def test_session_isolation_forgery_and_private_storage(app, client):
    course = demo(client)
    assert "_questions" not in course and "_attempts" not in course
    assert "text" not in course["materials"][0]
    with TestClient(app) as other:
        assert other.get("/api/bootstrap").status_code == 401
        other.cookies.set(COOKIE, "forged-session")
        assert other.get("/api/bootstrap").status_code == 401
        other.post("/api/session", json={"name": "Grace"})
        assert other.get("/api/bootstrap").json()["courses"][0]["id"] != course["id"]
        assert other.post(f"/api/courses/{course['id']}/sessions", json={}).status_code == 404
        mid = course["materials"][0]["id"]
        assert other.get(f"/api/courses/{course['id']}/materials/{mid}").status_code == 404
    cookie = client.cookies.get(COOKIE)
    with TestClient(create_app(app.state.store.path)) as restarted:
        restarted.cookies.set(COOKIE, cookie)
        assert restarted.get("/api/bootstrap").json()["courses"][0]["id"] == course["id"]


def test_csrf_rejected_and_same_origin_accepted(client):
    response = client.post("/api/courses", json={"name": "Hijack"}, headers={"origin": "https://malicious.invalid"})
    assert response.status_code == 403
    response = client.post("/api/courses", json={"name": "Allowed"}, headers={"origin": "http://testserver"})
    assert response.status_code == 200


def test_upload_lifecycle_missing_key_and_failed_parse(client):
    course = new_course(client)
    base = f"/api/courses/{course['id']}"
    uploaded = client.post(base + "/materials", files={"file": ("../../notes.md", b"# Algebra\nA group has an associative operation and inverses.", "text/markdown")}, data={"kind": "lecture"})
    assert uploaded.status_code == 200
    material = uploaded.json()
    assert material["name"] == "notes.md" and material["status"] == "ready"
    assert "text" not in material
    fetched = client.get(base + "/materials/" + material["id"]).json()
    assert "associative" in fetched["text"]
    before = next(c for c in client.get("/api/bootstrap").json()["courses"] if c["id"] == course["id"])
    assert before["topics"] == []  # Upload does not manufacture analysis.
    no_key = client.post(base + "/analyze")
    assert no_key.status_code == 503 and "OPENAI_API_KEY" in no_key.json()["detail"]
    invalid = client.post(base + "/materials", files={"file": ("broken.pdf", b"not a PDF", "application/pdf")})
    assert invalid.status_code == 200 and invalid.json()["status"] == "failed"
    assert invalid.json()["error"]
    unsupported = client.post(base + "/materials", files={"file": ("run.exe", b"something")})
    assert unsupported.status_code == 415
    assert client.delete(base + "/materials/" + material["id"]).json() == {"ok": True}
    assert client.get(base + "/materials/" + material["id"]).status_code == 404


def test_docx_and_pptx_extracts_actual_text(client):
    from docx import Document
    document = Document()
    document.add_paragraph("Bayes theorem relates conditional probabilities.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Prior"
    table.cell(0, 1).text = "0.2"
    data = io.BytesIO()
    document.save(data)
    course = new_course(client)
    base = f"/api/courses/{course['id']}"
    response = client.post(base + "/materials", files={"file": ("bayes.docx", data.getvalue())})
    assert response.json()["status"] == "ready"
    text = client.get(base + "/materials/" + response.json()["id"]).json()["text"]
    assert "Bayes theorem" in text and "Prior | 0.2" in text
    slides = io.BytesIO()
    with zipfile.ZipFile(slides, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", '<p:sld xmlns:p="urn:slide" xmlns:a="urn:paragraph"><a:t>Conditional probability</a:t></p:sld>')
    response = client.post(base + "/materials", files={"file": ("lesson.pptx", slides.getvalue())})
    assert response.json()["status"] == "ready" and response.json()["pages"] == 1


def test_practice_answer_is_private_and_mastery_updates_once(client):
    course = demo(client)
    topic = next(t for t in course["topics"] if "Recall" in t["title"])
    base = f"/api/courses/{course['id']}"
    response = client.post(base + "/questions", json={"topicId": topic["id"]})
    assert response.status_code == 200, response.text
    question = response.json()
    assert question["isDemo"] is True and question["sourceIds"]
    assert question["dna"]["reasoning"] == []
    assert not {"solution", "rubric", "_solution", "_rubric", "modelAnswer", "demoCorrectOption", "demoKey", "reviewDetails"}.intersection(question)
    rejected = client.post(base + f"/questions/{question['id']}/answers", json={"answer": "I do not know how to calculate this."})
    assert rejected.status_code == 422  # Demo only grades explicit choices; never fakes free-text grading.
    body = {"answer": "A"}
    first = client.post(base + f"/questions/{question['id']}/answers", json=body)
    assert first.status_code == 200, first.text
    feedback = first.json()
    assert feedback["maxScore"] > 0 and 0 <= feedback["score"] <= feedback["maxScore"]
    assert feedback["masteryBefore"] == topic["mastery"]
    assert feedback["masteryAfter"] != feedback["masteryBefore"]
    second = client.post(base + f"/questions/{question['id']}/answers", json=body)
    assert second.json() == feedback
    saved_topic = next(t for t in demo(client)["topics"] if t["id"] == topic["id"])
    assert saved_topic["mastery"] == feedback["masteryAfter"]


def test_chat_and_revision_pack_exports(client):
    course = demo(client)
    base = f"/api/courses/{course['id']}"
    session = client.post(base + "/sessions", json={}).json()
    chat = client.post(base + "/chat", json={"sessionId": session["id"], "message": "Explain recall", "language": "en"})
    assert chat.status_code == 200, chat.text
    assert chat.json()["isDemo"] and chat.json()["sourceIds"]
    generated = client.post(base + "/packs", json={"topicIds": [course["topics"][0]["id"]], "format": "notes", "detail": "quick", "language": "zh"})
    assert generated.status_code == 200, generated.text
    pack = generated.json()
    assert pack["isDemo"] and pack["content"]
    edited_content = "## 复习笔记\nPrecision = TP / (TP + FP).\n\n" + "Long line with useful details. " * 50
    edited = client.put(base + "/packs/" + pack["id"], json={"content": edited_content})
    assert edited.json()["content"] == edited_content
    for fmt, prefix in [("md", b"#"), ("pdf", b"%PDF"), ("docx", b"PK")]:
        response = client.get(base + f"/packs/{pack['id']}/export?format={fmt}")
        assert response.status_code == 200 and response.content.startswith(prefix)
        assert "attachment" in response.headers["content-disposition"]
    assert client.get(base + f"/packs/{pack['id']}/export?format=exe").status_code == 422


def test_chat_persists_route_metadata_and_switches_intent_per_message(client):
    course = demo(client)
    base = f"/api/courses/{course['id']}"
    session = client.post(base + "/sessions", json={}).json()
    generated = client.post(base + "/chat", json={"sessionId": session["id"], "message": "生成一份模拟考试卷", "language": "zh"})
    assert generated.status_code == 200
    assert generated.json()["route"] == "exam_generation"
    assert generated.json()["taskIntent"] == "generate_mock_exam"
    analyzed = client.post(base + "/chat", json={"sessionId": session["id"], "message": "总结这份试卷的高频考点", "language": "zh"})
    assert analyzed.status_code == 200
    assert analyzed.json()["route"] == "review_or_analysis"
    saved = next(item for item in demo(client)["sessions"] if item["id"] == session["id"])
    assistant_messages = [item for item in saved["messages"] if item["role"] == "assistant"]
    assert [item["route"] for item in assistant_messages] == ["exam_generation", "review_or_analysis"]


def test_share_snapshot_cross_browser_and_private_fields(app, client):
    course = demo(client)
    base = f"/api/courses/{course['id']}"
    share = client.post(base + "/share").json()
    assert "?share=" in share["url"]
    with TestClient(app) as visitor:
        response = visitor.get("/api/shared/" + share["token"])
        assert response.status_code == 200
        snapshot = response.json()
        assert snapshot["name"] == course["name"]
        forbidden = {"mastery", "readiness", "sessions", "messages", "answers", "_attempts", "_questions", "materials", "priority", "reasoning"}
        def inspect(value):
            if isinstance(value, dict):
                assert not forbidden.intersection(value)
                for subvalue in value.values(): inspect(subvalue)
            if isinstance(value, list):
                for item in value: inspect(item)
        inspect(snapshot)
        assert visitor.post("/api/shared/" + share["token"] + "/comments", json={"body": "Helpful"}).status_code == 401
        visitor.post("/api/session", json={"name": "Grace"})
        comment = visitor.post("/api/shared/" + share["token"] + "/comments", json={"body": "Let's study recall tomorrow."})
        assert comment.status_code == 200 and comment.json()["author"] == "Grace"
        assert len(client.get("/api/shared/" + share["token"]).json()["comments"]) == 1
        assert visitor.post(base + "/sessions", json={}).status_code == 404


def test_demo_upload_is_explicit_and_validation_is_human_readable(client):
    course = demo(client)
    response = client.post(f"/api/courses/{course['id']}/materials", files={"file": ("lecture.txt", b"real notes")})
    assert response.status_code == 409 and "own course" in response.json()["detail"]
    invalid = client.post("/api/courses", json={"name": ""})
    assert invalid.status_code == 422 and isinstance(invalid.json()["detail"], str)


def test_exam_dna_analysis_is_persisted_and_owner_scoped(app, client):
    course = demo(client)
    base = f"/api/courses/{course['id']}"
    assert client.get(base + "/exam-analysis").json() is None
    missing = client.post(base + "/exam-analysis", json={"language": "en"})
    assert missing.status_code == 200, missing.text
    analysis = missing.json()
    assert analysis["isDemo"] is True and analysis["sampleSize"] == 4
    assert sum(row["percentage"] for row in analysis["patterns"]) == 100
    assert analysis["yearRange"] == {"start": 2024, "end": 2026}
    assert all(item["sourceIds"] for item in analysis["questions"])
    assert client.get(base + "/exam-analysis").json() == analysis
    assert demo(client)["examAnalysis"] == analysis
    share = client.post(base + "/share").json()
    assert "examAnalysis" not in client.get("/api/shared/" + share["token"]).json()
    with TestClient(app) as other:
        assert other.get(base + "/exam-analysis").status_code == 401
        assert other.post(base + "/exam-analysis").status_code == 401
        other.post("/api/session", json={"name": "Other"})
        assert other.get(base + "/exam-analysis").status_code == 404
        assert other.post(base + "/exam-analysis").status_code == 404


def test_exam_analysis_requires_past_paper(client):
    course = new_course(client)
    response = client.post(f"/api/courses/{course['id']}/exam-analysis")
    assert response.status_code == 422 and "past paper" in response.json()["detail"].lower()
    base = f"/api/courses/{course['id']}"
    client.post(base + "/materials", files={"file": ("exam.txt", b"Q1 (5 marks): Explain probability.")}, data={"kind": "past_paper"})
    missing_key = client.post(base + "/exam-analysis", json={"language": "en"})
    assert missing_key.status_code == 503 and "OPENAI_API_KEY" in missing_key.json()["detail"]
    bad_language = client.post(base + "/exam-analysis", json={"language": "invalid"})
    assert bad_language.status_code == 422 and isinstance(bad_language.json()["detail"], str)


def test_removing_evidence_invalidates_exam_analysis_and_packs(app, client):
    course = new_course(client)
    base = f"/api/courses/{course['id']}"
    material = client.post(base + "/materials", files={"file": ("exam.txt", b"Q1: Explain probability.")}, data={"kind": "past_paper"}).json()
    owner = client.get("/api/bootstrap").json()["user"]["id"]
    saved = app.state.store.course(owner, course["id"])
    saved.update(examAnalysis={"sampleSize": 1}, topics=[{"id": "stale", "mastery": 70}], packs=[{"id": "stale", "content": "Evidence removed"}])
    app.state.store.save_course(owner, saved)
    assert client.delete(base + "/materials/" + material["id"]).status_code == 200
    assert client.get(base + "/exam-analysis").json() is None
    current = next(item for item in client.get("/api/bootstrap").json()["courses"] if item["id"] == course["id"])
    assert current["topics"] == [] and current["packs"] == []


def test_static_frontend_preserves_api_routes_and_errors(tmp_path):
    dist = tmp_path / "dist"
    dist.mkdir()
    (dist / "index.html").write_text("<!doctype html><main>Exam Radar</main>", encoding="utf-8")
    (dist / "assets").mkdir()
    (dist / "assets" / "app.js").write_text("window.examRadar = true;", encoding="utf-8")
    with TestClient(create_app(tmp_path / "test.sqlite3", static_dir=dist)) as browser:
        assert "<main>Exam Radar</main>" in browser.get("/").text
        assert "<main>Exam Radar</main>" in browser.get("/?share=demo-token").text
        assert browser.get("/assets/app.js").status_code == 200
        assert browser.get("/api/health").json()["status"] == "ok"
        unknown = browser.get("/api/missing")
        assert unknown.status_code == 404 and unknown.json()["detail"] == "API endpoint not found."
        assert browser.get("/api/bootstrap").status_code == 401


def test_concurrent_uploads_do_not_drop_material_records(client):
    course = new_course(client)
    base = f"/api/courses/{course['id']}"

    def upload(index):
        return client.post(base + "/materials", files={"file": (f"notes-{index}.txt", f"Learning outcome {index}: understand vectors.".encode())})

    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(upload, range(8)))
    assert all(response.status_code == 200 for response in responses)
    saved_course = next(c for c in client.get("/api/bootstrap").json()["courses"] if c["id"] == course["id"])
    assert len(saved_course["materials"]) == 8
    assert len({m["id"] for m in saved_course["materials"]}) == 8
