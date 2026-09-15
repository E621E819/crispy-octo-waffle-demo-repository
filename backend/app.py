"""Cookie-isolated local MVP API. Run: python -m uvicorn backend.app:app."""
from __future__ import annotations

import asyncio
import io
import logging
import os
import re
import secrets
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Literal
from urllib.parse import quote, urlsplit
from xml.etree import ElementTree

from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, Response, RedirectResponse
from fastapi.staticfiles import StaticFiles
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from pydantic import BaseModel, Field

from .export import export_pack
from .seed import demo_course, readiness, uid
from .store import Store
from .auth import AUTH_COOKIE, BROWSER_COOKIE, SESSION_TTL, STATE_TTL, authorization_url, constant_time_equal, random_id, reviewer_enabled, safe_redirect, zhihu_config

load_dotenv(Path(__file__).resolve().parent.parent / ".env")
logger = logging.getLogger("exam_radar")
COOKIE = "exam_radar_session"
MAX_UPLOAD = 15 * 1024 * 1024
MAX_TEXT = 1_000_000


class SessionBody(BaseModel):
    name: str | None = Field(default=None, max_length=80)

class ReviewerLoginBody(BaseModel):
    username: str = Field(min_length=1, max_length=160)
    password: str = Field(min_length=1, max_length=256)


class CourseBody(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    code: str = Field(default="", max_length=40)
    examDate: date | None = None


class ConversationBody(BaseModel):
    title: str = Field(default="New study session", min_length=1, max_length=120)


class ChatBody(BaseModel):
    sessionId: str
    message: str = Field(min_length=1, max_length=12000)
    language: Literal["en", "zh"] = "zh"


class QuestionBody(BaseModel):
    topicId: str
    language: Literal["en", "zh"] = "zh"


class ExamAnalysisBody(BaseModel):
    language: Literal["en", "zh"] = "zh"


class AnswerBody(BaseModel):
    answer: str = Field(min_length=1, max_length=12000)
    language: Literal["en", "zh"] = "zh"


class PackBody(BaseModel):
    topicIds: list[str] = Field(min_length=1, max_length=40)
    format: Literal["notes", "formulas", "flashcards"] = "notes"
    detail: Literal["quick", "standard", "full"] = "standard"
    language: Literal["en", "zh"] = "zh"


class PackEditBody(BaseModel):
    content: str = Field(max_length=200000)


class CommentBody(BaseModel):
    body: str = Field(min_length=1, max_length=3000)


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def public_material(material: dict) -> dict:
    return {k: v for k, v in material.items() if k in {"id", "name", "kind", "status", "pages", "excerpt", "error", "isDemo"}}


def public_course(course: dict) -> dict:
    result = {k: course.get(k) for k in ("id", "name", "code", "examDate", "isDemo", "topics", "sessions", "readiness", "examAnalysis")}
    result["materials"] = [public_material(m) for m in course.get("materials", [])]
    result["packs"] = course.get("packs", [])
    return result


def public_question(question: dict) -> dict:
    fields = {"id", "topicId", "prompt", "options", "marks", "difficulty", "dna", "checks", "sourceIds", "isDemo"}
    result = {k: v for k, v in question.items() if k in fields}
    # DNA solution milestones are useful for grading but may give away the answer.
    # Keep them private together with the rubric until the answer is submitted.
    dna = question.get("dna", {})
    result["dna"] = {k: dna[k] for k in ("concept", "cognitiveLevel", "transformation") if k in dna}
    result["dna"]["reasoning"] = []
    return result


def extract_document(data: bytes, filename: str) -> tuple[str, int]:
    extension = Path(filename).suffix.lower()
    if extension not in {".txt", ".md", ".pdf", ".docx", ".pptx"}:
        raise HTTPException(415, "Supported file types: PDF, DOCX, PPTX, TXT and Markdown.")
    if not data:
        raise ValueError("This file is empty.")
    if extension in {".txt", ".md"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Save text files with UTF-8 encoding and upload again.")
        return f"[Page 1]\n{text}", 1
    if extension == ".pdf":
        import fitz
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise ValueError("Password-protected PDFs are not supported. Upload an unlocked copy.")
            if len(document) > 500:
                raise ValueError("Please split PDFs longer than 500 pages.")
            pages = [f"[Page {i + 1}]\n{page.get_text()}" for i, page in enumerate(document)]
            content = "\n\n".join(pages)
            # Page labels alone are not source text; scanned PDFs require OCR upstream.
            if not any(page.get_text().strip() for page in document):
                raise ValueError("This PDF contains no selectable text. Run OCR or upload a text-based document.")
            return content, len(document)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise ValueError("Expanded document is too large. Split it into smaller files.")
        if extension == ".docx":
            from docx import Document
            document = Document(io.BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "[Page 1 · DOCX text; original pagination unavailable]\n" + "\n".join(paragraphs), 1
        slides = sorted((entry.filename for entry in entries if re.fullmatch(r"ppt/slides/slide\d+\.xml", entry.filename)), key=lambda name: int(re.search(r"slide(\d+)", name).group(1)))
        if len(slides) > 500:
            raise ValueError("Please split presentations longer than 500 slides.")
        pages = []
        for number, slide in enumerate(slides, 1):
            root = ElementTree.fromstring(archive.read(slide))
            texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
            pages.append(f"[Page {number} · Slide {number}]\n" + "\n".join(texts))
        return "\n\n".join(pages), len(slides)


def create_app(db_path: str | Path | None = None, static_dir: str | Path | None = None) -> FastAPI:
    application = FastAPI(title="Exam Radar", version="0.1.0")
    database = Store(db_path)
    application.state.store = database
    application.state.course_locks = {}

    @application.exception_handler(RequestValidationError)
    async def validation_error(request: Request, exc: RequestValidationError):
        issue = exc.errors()[0]
        location = ".".join(str(part) for part in issue.get("loc", ()) if part not in {"body", "query"})
        message = issue.get("msg", "Invalid request.")
        return JSONResponse({"detail": f"{location}: {message}" if location else message}, status_code=422)
    configured_secret = os.getenv("EXAM_RADAR_SESSION_SECRET")
    secret_path = database.path.parent / ".session-secret"
    if configured_secret:
        secret = configured_secret
    else:
        # Persist a randomly generated signing key so sessions survive local restarts.
        try:
            with secret_path.open("x", encoding="utf-8") as output:
                output.write(secrets.token_urlsafe(48))
        except FileExistsError:
            pass
        secret = secret_path.read_text(encoding="utf-8")
    signer = URLSafeTimedSerializer(secret, salt="exam-radar-browser-session-v1")

    def mode():
        configured = os.getenv("AI_MODE", "hybrid").strip().lower()
        if configured not in {"hybrid", "local", "live", "demo"}:
            configured = "hybrid"
        if configured == "live" and not os.getenv("OPENAI_API_KEY", "").strip():
            return "hybrid"
        return configured

    def user_for(request: Request, required: bool = True):
        auth_raw = request.cookies.get(AUTH_COOKIE)
        if auth_raw:
            item = database.auth_session(auth_raw)
            if item and not item.get("revoked") and float(item.get("expiresAt", 0)) > __import__('time').time():
                user = database.user(item["userId"])
                if user:
                    return user
        raw = request.cookies.get(COOKIE)
        user = None
        if raw:
            try:
                identifier = signer.loads(raw, max_age=60 * 60 * 24 * 30)
                user = database.user(identifier)
            except (BadSignature, SignatureExpired, TypeError):
                pass
        if required and not user:
            raise HTTPException(401, "Start a local session before continuing.")
        return user

    def course_for(request: Request, course_id: str):
        user = user_for(request)
        course = database.course(user["id"], course_id)
        if not course:
            raise HTTPException(404, "Course not found in your workspace.")
        return user, course

    def lock_for(course_id: str):
        return application.state.course_locks.setdefault(course_id, asyncio.Lock())

    def save_course(user: dict, course: dict):
        course["readiness"] = readiness(course.get("topics", []))
        database.save_course(user["id"], course)

    async def ai_call(course: dict, method: str, **kwargs):
        from .ai import AIService, AIServiceError
        if not course.get("isDemo") and mode() == "demo":
            raise HTTPException(503, "当前为演示模式。请将 AI_MODE 设置为 hybrid 并启动 Ollama，或配置 OPENAI_API_KEY 后重启服务。")
        if course.get("isDemo") and any(not item.get("isDemo") for item in course.get("materials", [])):
            # Uploads change the course to real mode; protect old data migrated from earlier builds.
            raise HTTPException(409, "Analyze uploaded documents in a new course so sample data is not treated as real analysis.")
        try:
            return await getattr(AIService(), method)(course=course, **kwargs)
        except AIServiceError as exc:
            raise HTTPException(getattr(exc, "status_code", 502), str(exc)) from exc
        except Exception as exc:
            logger.exception("AI request failed")
            raise HTTPException(502, "The learning service could not finish this request. Retry, or check the server log and API configuration.") from exc

    @application.middleware("http")
    async def protect_origin(request: Request, call_next):
        if request.method not in {"GET", "HEAD", "OPTIONS"}:
            origin = request.headers.get("origin")
            if origin and urlsplit(origin).hostname != request.url.hostname:
                return JSONResponse({"detail": "Cross-site requests are not allowed."}, status_code=403)
            if request.headers.get("sec-fetch-site") == "cross-site":
                return JSONResponse({"detail": "Cross-site requests are not allowed."}, status_code=403)
        response = await call_next(request)
        if request.url.path.startswith("/api/"):
            response.headers["Cache-Control"] = "no-store"
            response.headers["X-Content-Type-Options"] = "nosniff"
        return response

    @application.get("/api/health")
    def health():
        return {"status": "ok", "mode": mode(), "model": "gpt-5", "localModel": os.getenv("LOCAL_MODEL", "qwen2.5:7b")}

    @application.get("/api/auth/config")
    def auth_config():
        return zhihu_config()

    @application.post("/api/auth/reviewer")
    def reviewer_login(request: Request, body: ReviewerLoginBody):
        import time
        configured_user = os.getenv("EXAM_RADAR_REVIEWER_USERNAME", "")
        configured_password = os.getenv("EXAM_RADAR_REVIEWER_PASSWORD", "")
        key = f"reviewer:{request.client.host if request.client else 'unknown'}"
        limited = database.rate_limit(key, time.time(), 300) > 10
        valid = reviewer_enabled() and constant_time_equal(body.username, configured_user) and constant_time_equal(body.password, configured_password)
        if limited or not valid:
            raise HTTPException(401, "Invalid reviewer credentials.")
        # Each browser gets its own user/workspace, while the identity is shared for stats.
        browser = request.cookies.get(BROWSER_COOKIE) or random_id("browser")
        anon = user_for(request, required=False)
        if not anon:
            anon = {"id": random_id("user"), "name": body.username}
            database.save_user(anon); database.save_course(anon["id"], demo_course())
        identity = database.auth_identity("reviewer", configured_user)
        if not identity:
            identity = {"id": random_id("identity"), "provider": "reviewer", "subject": configured_user, "displayName": configured_user}
            database.save_auth_identity(identity)
        existing_sid = request.cookies.get(AUTH_COOKIE)
        existing = database.auth_session(existing_sid) if existing_sid else None
        if existing and not existing.get("revoked") and float(existing.get("expiresAt", 0)) > time.time() and existing.get("identityId") == identity["id"]:
            return {"user": database.user(existing["userId"]), "provider": "reviewer"}
        if existing and existing.get("id"): database.revoke_auth_session(existing["id"])
        sid = random_id("auth")
        database.save_auth_session({"id": sid, "identityId": identity["id"], "userId": anon["id"], "provider": "reviewer", "expiresAt": time.time()+SESSION_TTL, "revoked": False})
        database.record_login({"id": random_id("login"), "identityId": identity["id"], "provider": "reviewer", "createdAt": now()})
        response = JSONResponse({"user": anon, "provider": "reviewer"})
        response.set_cookie(AUTH_COOKIE, sid, httponly=True, secure=request.url.scheme == "https", samesite="lax", max_age=SESSION_TTL)
        response.set_cookie(BROWSER_COOKIE, browser, httponly=True, secure=request.url.scheme == "https", samesite="lax", max_age=SESSION_TTL)
        return response

    @application.get("/api/auth/me")
    def auth_me(request: Request):
        item = database.auth_session(request.cookies.get(AUTH_COOKIE, "")) if request.cookies.get(AUTH_COOKIE) else None
        if not item or item.get("revoked") or float(item.get("expiresAt", 0)) <= __import__('time').time():
            return {"authenticated": False, "user": None, "provider": "anonymous"}
        return {"authenticated": True, "user": database.user(item["userId"]), "provider": item.get("provider")}

    @application.post("/api/auth/logout")
    def auth_logout(request: Request):
        sid = request.cookies.get(AUTH_COOKIE)
        if sid: database.revoke_auth_session(sid)
        response = JSONResponse({"ok": True})
        response.delete_cookie(AUTH_COOKIE)
        return response

    @application.get("/api/admin/login-stats")
    def login_stats(request: Request):
        expected = os.getenv("EXAM_RADAR_ADMIN_TOKEN", "").strip()
        supplied = request.headers.get("authorization", "")
        token = supplied[7:].strip() if supplied.lower().startswith("bearer ") else ""
        if not expected or not constant_time_equal(token, expected): raise HTTPException(403, "Forbidden")
        return database.login_stats()

    @application.get("/api/auth/zhihu/start")
    def zhihu_start(request: Request):
        import time
        cfg = zhihu_config()
        if not cfg["zhihuEnabled"]: raise HTTPException(503, cfg["zhihuReason"])
        state = secrets.token_urlsafe(32); redirect = safe_redirect(request, os.getenv("ZHIHU_REDIRECT_URI", ""))
        browser_id = request.cookies.get(BROWSER_COOKIE) or random_id("browser")
        database.save_oauth_state({"state": state, "browserId": browser_id, "redirectUri": redirect, "expiresAt": time.time()+STATE_TTL})
        response = RedirectResponse(authorization_url(request, state, redirect), status_code=302)
        if not request.cookies.get(BROWSER_COOKIE): response.set_cookie(BROWSER_COOKIE, browser_id, httponly=True, secure=request.url.scheme == "https", samesite="lax", max_age=SESSION_TTL)
        return response

    @application.get("/api/auth/zhihu/callback")
    async def zhihu_callback(request: Request, code: str = "", state: str = "", error: str = ""):
        import time, httpx
        item = database.pop_oauth_state(state) if state else None
        browser = request.cookies.get(BROWSER_COOKIE, "")
        if not item or float(item.get("expiresAt", 0)) <= time.time() or not browser or not constant_time_equal(browser, item.get("browserId", "")):
            raise HTTPException(400, "Invalid or expired OAuth state.")
        if error or not code: raise HTTPException(400, "知乎登录未完成。")
        cfg = zhihu_config()
        if not cfg["zhihuEnabled"]: raise HTTPException(503, cfg["zhihuReason"])
        try:
            async with httpx.AsyncClient(timeout=15, follow_redirects=False) as client:
                token_resp = await client.post(os.getenv("ZHIHU_TOKEN_URL"), data={"grant_type":"authorization_code","code":code,"redirect_uri":item["redirectUri"],"client_id":os.getenv("ZHIHU_CLIENT_ID", ""),"client_secret":os.getenv("ZHIHU_CLIENT_SECRET", "")})
                token_resp.raise_for_status(); token_data = token_resp.json(); access = token_data.get("access_token", "")
                if not access: raise ValueError("missing token")
                user_resp = await client.get(os.getenv("ZHIHU_USERINFO_URL"), headers={"Authorization": f"Bearer {access}"})
                user_resp.raise_for_status(); profile = user_resp.json()
            subject = profile
            for part in (os.getenv("ZHIHU_USER_ID_FIELD", "id") or "id").split("."): subject = subject.get(part) if isinstance(subject, dict) else None
            if not subject: raise ValueError("missing subject")
        except Exception as exc:
            logger.warning("Zhihu OAuth exchange failed: %s", type(exc).__name__)
            raise HTTPException(502, "知乎登录暂时不可用，请稍后重试。") from exc
        user = user_for(request, required=False)
        if not user:
            user = {"id": random_id("user"), "name": str(profile.get("name") or "知乎用户")[:80]}; database.save_user(user); database.save_course(user["id"], demo_course())
        identity = database.auth_identity("zhihu", str(subject)) or {"id": random_id("identity"), "provider":"zhihu", "subject":str(subject), "displayName":str(profile.get("name") or "知乎用户")[:80]}
        database.save_auth_identity(identity); sid = random_id("auth")
        database.save_auth_session({"id":sid,"identityId":identity["id"],"userId":user["id"],"provider":"zhihu","expiresAt":time.time()+SESSION_TTL,"revoked":False})
        database.record_login({"id":random_id("login"),"identityId":identity["id"],"provider":"zhihu","createdAt":now()})
        response = RedirectResponse(str(request.base_url).rstrip("/") + "/", status_code=302)
        response.set_cookie(AUTH_COOKIE, sid, httponly=True, secure=request.url.scheme == "https", samesite="lax", max_age=SESSION_TTL)
        return response

    @application.post("/api/session")
    def session(request: Request, body: SessionBody | None = None):
        user = user_for(request, required=False)
        name = body.name.strip() if body and body.name else ""
        if not user:
            user = {"id": uid("user"), "name": name or "Study partner"}
            database.save_user(user)
            database.save_course(user["id"], demo_course())
        elif name and name != user["name"] and not request.cookies.get(AUTH_COOKIE):
            user["name"] = name
            database.save_user(user)
        response = JSONResponse({"user": user, "mode": mode(), "model": "gpt-5"})
        response.set_cookie(COOKIE, signer.dumps(user["id"]), httponly=True, samesite="lax", secure=request.url.scheme == "https", max_age=60 * 60 * 24 * 30)
        return response

    @application.get("/api/bootstrap")
    def bootstrap(request: Request):
        user = user_for(request)
        return {"user": user, "mode": mode(), "model": "gpt-5", "courses": [public_course(course) for course in database.courses(user["id"])]}

    @application.post("/api/courses")
    def create_course(request: Request, body: CourseBody):
        user = user_for(request)
        if not body.name.strip():
            raise HTTPException(422, "Give this course a name.")
        course = {"id": uid("course"), "name": body.name.strip(), "code": body.code.strip(), "examDate": body.examDate.isoformat() if body.examDate else None, "isDemo": False, "topics": [], "materials": [], "sessions": [], "readiness": None, "packs": [], "_questions": [], "_attempts": []}
        save_course(user, course)
        return public_course(course)

    @application.post("/api/courses/{course_id}/sessions")
    async def create_conversation(request: Request, course_id: str, body: ConversationBody | None = None):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            item = {"id": uid("session"), "title": body.title.strip() if body else "New study session", "messages": []}
            course["sessions"].append(item)
            save_course(user, course)
            return item

    @application.post("/api/courses/{course_id}/chat")
    async def chat(request: Request, course_id: str, body: ChatBody):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            item = next((session for session in course["sessions"] if session["id"] == body.sessionId), None)
            if item is None:
                raise HTTPException(404, "Study session not found.")
            if not body.message.strip():
                raise HTTPException(422, "Write a question before sending.")
            result = await ai_call(course, "chat", materials=course["materials"], history=item["messages"], message=body.message.strip(), language=body.language)
            message = {"id": uid("message"), "role": "assistant", "content": result["content"],
                       "sourceIds": result.get("sourceIds", []), "isDemo": bool(course["isDemo"]),
                       "route": result.get("route", "general_chat"),
                       "sourceType": result.get("sourceType", "unknown"),
                       "taskIntent": result.get("taskIntent", "general_chat")}
            item["messages"].extend([{"id": uid("message"), "role": "user", "content": body.message.strip()}, message])
            if item["title"] == "New study session":
                item["title"] = body.message.strip()[:60]
            save_course(user, course)
            return message

    @application.post("/api/courses/{course_id}/materials")
    async def upload_material(request: Request, course_id: str, file: UploadFile = File(...), kind: str = Form("lecture")):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            if course["isDemo"]:
                raise HTTPException(409, "Create your own course to upload documents. The sample course stays available for the demo.")
            if len(course["materials"]) >= 40:
                raise HTTPException(422, "This course already has 40 files. Remove an unused file first.")
            if len(kind) > 40 or not kind.strip():
                raise HTTPException(422, "Choose a document type.")
            filename = re.split(r"[/\\]", file.filename or "document")[-1][:180]
            if Path(filename).suffix.lower() not in {".pdf", ".docx", ".pptx", ".txt", ".md"}:
                raise HTTPException(415, "Supported file types: PDF, DOCX, PPTX, TXT and Markdown.")
            data = await file.read(MAX_UPLOAD + 1)
            await file.close()
            if len(data) > MAX_UPLOAD:
                raise HTTPException(413, "Files must be 15 MB or smaller.")
            material = {"id": uid("material"), "name": filename, "kind": kind.strip(), "status": "processing", "pages": 0, "excerpt": "", "isDemo": False, "text": ""}
            try:
                text, pages = await asyncio.to_thread(extract_document, data, filename)
                if len(text) > MAX_TEXT:
                    raise ValueError("Extracted text exceeds one million characters. Split the document and upload again.")
                meaningful = re.sub(r"\[Page[^\]]*\]", "", text).strip()
                if not meaningful:
                    raise ValueError("No readable text was found. Try an OCR or text export of this file.")
                material.update(status="ready", pages=pages, text=text, excerpt=text[:320])
            except HTTPException:
                raise
            except Exception as exc:
                error = str(exc) if isinstance(exc, ValueError) else "The file could not be parsed. Verify its format, or export it as PDF or UTF-8 text."
                material.update(status="failed", error=error)
            course["materials"].append(material)
            save_course(user, course)
            return public_material(material)

    @application.get("/api/courses/{course_id}/materials/{material_id}")
    def get_material(request: Request, course_id: str, material_id: str):
        _, course = course_for(request, course_id)
        material = next((m for m in course["materials"] if m["id"] == material_id), None)
        if not material:
            raise HTTPException(404, "Source document not found.")
        return {"material": public_material(material), "text": material.get("text", "")}

    @application.delete("/api/courses/{course_id}/materials/{material_id}")
    async def remove_material(request: Request, course_id: str, material_id: str):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            if course["isDemo"]:
                raise HTTPException(409, "Sample documents support the demo. Create your own course to manage your documents.")
            if not any(m["id"] == material_id for m in course["materials"]):
                raise HTTPException(404, "Source document not found.")
            course["materials"] = [m for m in course["materials"] if m["id"] != material_id]
            # Invalidate analysis whose evidence was removed; never keep stale priorities.
            course["topics"] = []
            course["_questions"] = []
            course["examAnalysis"] = None
            course["packs"] = []
            save_course(user, course)
            return {"ok": True}

    @application.post("/api/courses/{course_id}/analyze")
    async def analyze(request: Request, course_id: str):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            materials = [m for m in course["materials"] if m["status"] == "ready"]
            if not materials:
                raise HTTPException(422, "Upload at least one readable source document before analyzing.")
            topics = await ai_call(course, "analyze", materials=materials)
            course["topics"] = topics
            course["_questions"] = []
            save_course(user, course)
            return public_course(course)

    @application.post("/api/courses/{course_id}/exam-analysis")
    async def create_exam_analysis(request: Request, course_id: str, body: ExamAnalysisBody | None = None):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            materials = [m for m in course["materials"] if m.get("status") == "ready"]
            if not materials:
                raise HTTPException(422, "Upload a readable past paper before extracting Exam DNA.")
            language = body.language if body else "zh"
            analysis = await ai_call(course, "exam_analysis", materials=materials, language=language)
            course["examAnalysis"] = analysis
            save_course(user, course)
            return analysis

    @application.get("/api/courses/{course_id}/exam-analysis")
    def get_exam_analysis(request: Request, course_id: str):
        _, course = course_for(request, course_id)
        return course.get("examAnalysis")

    @application.post("/api/courses/{course_id}/questions")
    async def create_question(request: Request, course_id: str, body: QuestionBody):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            topic = next((t for t in course["topics"] if t["id"] == body.topicId), None)
            if not topic:
                raise HTTPException(404, "Topic not found. Analyze course documents first.")
            question = await ai_call(course, "generate_question", topic=topic, materials=course["materials"], language=body.language)
            question.update(id=uid("question"), topicId=topic["id"], isDemo=bool(course["isDemo"]))
            course.setdefault("_questions", []).append(question)
            save_course(user, course)
            return public_question(question)

    @application.post("/api/courses/{course_id}/questions/{question_id}/answers")
    async def submit_answer(request: Request, course_id: str, question_id: str, body: AnswerBody):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            question = next((q for q in course.get("_questions", []) if q["id"] == question_id), None)
            if question is None:
                raise HTTPException(404, "Practice question not found. Generate another question.")
            if not body.answer.strip():
                raise HTTPException(422, "Write an answer before submitting.")
            previous = next((a for a in course.get("_attempts", []) if a["questionId"] == question_id), None)
            if previous:
                # Browser retries must not award repeated mastery updates.
                return previous["feedback"]
            topic = next((t for t in course["topics"] if t["id"] == question["topicId"]), None)
            if topic is None:
                raise HTTPException(409, "The course map has changed. Generate a new practice question.")
            feedback = await ai_call(course, "grade", question=question, answer=body.answer.strip(), materials=course["materials"], language=body.language)
            from .services import next_topic_id, priority_score, update_mastery
            before = topic.get("mastery")
            after = update_mastery(before, feedback["score"], feedback["maxScore"])
            topic["mastery"] = after
            topic["priority"] = priority_score(topic)
            feedback.update(masteryBefore=before, masteryAfter=after, nextTopicId=next_topic_id(course["topics"], topic["id"]), isDemo=bool(course["isDemo"]))
            course.setdefault("_attempts", []).append({"id": uid("attempt"), "questionId": question_id, "answer": body.answer.strip(), "feedback": feedback, "createdAt": now()})
            save_course(user, course)
            return feedback

    @application.post("/api/courses/{course_id}/packs")
    async def create_pack(request: Request, course_id: str, body: PackBody):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            topics = [t for t in course["topics"] if t["id"] in body.topicIds]
            if len(topics) != len(set(body.topicIds)):
                raise HTTPException(404, "A selected topic is no longer available. Refresh the course map.")
            pack = await ai_call(course, "revision_pack", topics=topics, materials=course["materials"], format=body.format, detail=body.detail, language=body.language)
            pack.update(id=uid("pack"), createdAt=now(), isDemo=bool(course["isDemo"]))
            course.setdefault("packs", []).append(pack)
            save_course(user, course)
            return pack

    @application.get("/api/courses/{course_id}/packs")
    def get_packs(request: Request, course_id: str):
        return course_for(request, course_id)[1].get("packs", [])

    @application.put("/api/courses/{course_id}/packs/{pack_id}")
    async def edit_pack(request: Request, course_id: str, pack_id: str, body: PackEditBody):
        async with lock_for(course_id):
            user, course = course_for(request, course_id)
            pack = next((p for p in course.get("packs", []) if p["id"] == pack_id), None)
            if pack is None:
                raise HTTPException(404, "Revision pack not found.")
            pack["content"] = body.content
            save_course(user, course)
            return pack

    @application.get("/api/courses/{course_id}/packs/{pack_id}/export")
    def download_pack(request: Request, course_id: str, pack_id: str, format: Literal["docx", "pdf", "md"] = "md"):
        _, course = course_for(request, course_id)
        pack = next((p for p in course.get("packs", []) if p["id"] == pack_id), None)
        if pack is None:
            raise HTTPException(404, "Revision pack not found.")
        data, media_type = export_pack(pack, format)
        filename = re.sub(r"[^\w\- ]", "", pack["title"])[:80].strip() or "revision-pack"
        return Response(data, media_type=media_type, headers={"Content-Disposition": f"attachment; filename=revision-pack.{format}; filename*=UTF-8''{quote(filename + '.' + format)}"})

    @application.post("/api/courses/{course_id}/share")
    def create_share(request: Request, course_id: str):
        user, course = course_for(request, course_id)
        token = secrets.token_urlsafe(24)
        # Explicit allowlist prevents mastery, weighted priorities, reasoning about weaknesses,
        # private chat, attempts and source document contents leaking into a link snapshot.
        topics = [{k: t[k] for k in ("id", "title", "chapter", "summary") if k in t} for t in course["topics"]]
        packs = [{k: p[k] for k in ("id", "title", "content", "topicIds", "createdAt", "isDemo") if k in p} for p in course.get("packs", [])]
        snapshot = {"name": course["name"], "topics": topics, "packs": packs, "isDemo": course["isDemo"], "createdAt": now()}
        database.create_share(user["id"], course_id, token, snapshot, now())
        return {"token": token, "url": f"/?share={token}"}

    @application.get("/api/shared/{token}")
    def shared_course(token: str):
        snapshot = database.share(token)
        if snapshot is None:
            raise HTTPException(404, "This shared snapshot was not found.")
        return snapshot

    @application.post("/api/shared/{token}/comments")
    def add_comment(request: Request, token: str, body: CommentBody):
        user = user_for(request)
        if database.share(token) is None:
            raise HTTPException(404, "This shared snapshot was not found.")
        if not body.body.strip():
            raise HTTPException(422, "Write a comment before posting.")
        comment = {"id": uid("comment"), "author": user["name"], "body": body.body.strip(), "createdAt": now()}
        database.add_comment(token, user, comment)
        return comment

    @application.api_route("/api/{unmatched_path:path}", methods=["GET", "HEAD", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"])
    def unknown_api(unmatched_path: str):
        raise HTTPException(404, "API endpoint not found.")

    frontend_dist = Path(static_dir) if static_dir is not None else Path(__file__).resolve().parent.parent / "frontend" / "dist"
    if frontend_dist.is_dir():
        # API routes remain ordered first; query-based share links use index.html.
        application.mount("/", StaticFiles(directory=frontend_dist, html=True), name="frontend")

    return application


app = create_app()
