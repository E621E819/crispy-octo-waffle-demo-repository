"""In-memory revision pack downloads with Unicode support and no temp-file leaks."""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape


def export_pack(pack: dict, fmt: str) -> tuple[bytes, str]:
    title = pack.get("title", "Revision Pack")
    content = pack.get("content", "")
    if fmt == "md":
        return f"# {title}\n\n{content}\n".encode("utf-8"), "text/markdown; charset=utf-8"
    if fmt == "docx":
        from docx import Document
        from docx.shared import Pt
        document = Document()
        document.add_heading(title, 0)
        document.styles["Normal"].font.size = Pt(11)
        for line in content.splitlines():
            heading = re.match(r"^(#{1,6})\s+(.*)", line)
            if heading:
                document.add_heading(heading.group(2), min(heading.group(1).count("#"), 4))
            elif line.startswith(("- ", "* ")):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_LEFT
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.pagesizes import A4
        # Built-in CJK CID font supports bilingual notes without platform font assumptions.
        try:
            pdfmetrics.getFont("STSong-Light")
        except KeyError:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48, title=title)
        normal = ParagraphStyle("Notes", fontName="STSong-Light", fontSize=11, leading=17, textColor=colors.HexColor("#292925"), wordWrap="CJK", alignment=TA_LEFT)
        heading_style = ParagraphStyle("Heading", parent=normal, fontSize=16, leading=22, spaceBefore=12, spaceAfter=8)
        story = [Paragraph(escape(title), heading_style), Spacer(1, 12)]
        for line in content.splitlines():
            heading = re.match(r"^#{1,6}\s+(.*)", line)
            if heading:
                story.append(Paragraph(escape(heading.group(1)), heading_style))
            elif line:
                story.append(Paragraph(escape(line), normal))
            else:
                story.append(Spacer(1, 7))
        doc.build(story)
        return output.getvalue(), "application/pdf"
    raise ValueError("Choose md, docx or pdf for export.")
